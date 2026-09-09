"""Compiles a dossier (and optional match result) into a PDF or DOCX bundle.

Spec reference: MASTER_DEVELOPER_SPEC.pdf Section 3 defines ``qv export
--bundle <name> --format <pdf|docx>`` as compiling "tailored resume and
narrative dossier into formatted PDF." The spec's visual theme (Section 1)
is for the interactive HUD application, not the exported document, so —
per the Stage 1 blueprint — this module uses a plain, professional document
layout rather than importing the neon HUD palette into a resume, which
would be inappropriate for a document meant to be read by human recruiters.
"""

from __future__ import annotations

from pathlib import Path
from typing import Literal

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from questvector.core.models import Dossier, MatchResult
from questvector.exceptions import ExportError

ExportFormat = Literal["pdf", "docx"]
_SUPPORTED_FORMATS: frozenset[str] = frozenset({"pdf", "docx"})


def _validate_format(fmt: str) -> None:
    """Validate that an export format is supported.

    Args:
        fmt: The requested export format string.

    Raises:
        ExportError: If ``fmt`` is not one of :data:`_SUPPORTED_FORMATS`.
    """
    if fmt not in _SUPPORTED_FORMATS:
        raise ExportError(
            f"Unsupported export format: {fmt!r}",
            context={"supported_formats": sorted(_SUPPORTED_FORMATS)},
        )


def export_bundle(
    dossier: Dossier,
    output_path: Path,
    *,
    fmt: ExportFormat,
    match_result: MatchResult | None = None,
    narrative: str | None = None,
    title: str = "Career Dossier",
) -> Path:
    """Compile a dossier bundle into a PDF or DOCX file.

    Args:
        dossier: The parsed dossier to render.
        output_path: Destination file path. Parent directories must exist.
        fmt: ``"pdf"`` or ``"docx"``.
        match_result: Optional G-Force match result to include as a
            summary/gaps section.
        narrative: Optional LLM-generated narrative summary paragraph (see
            :mod:`questvector.core.llm`) to include at the top of the
            document.
        title: Document title.

    Returns:
        ``output_path``, on success.

    Raises:
        ExportError: If ``fmt`` is unsupported, or the underlying
            PDF/DOCX library raises while building the document.
    """
    _validate_format(fmt)
    try:
        if fmt == "pdf":
            _export_pdf(dossier, output_path, match_result=match_result, narrative=narrative, title=title)
        else:
            _export_docx(dossier, output_path, match_result=match_result, narrative=narrative, title=title)
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - deliberately broad: wrap any
        # third-party rendering-library failure into our own exception type
        # so callers only ever need to catch QuestVectorError subclasses.
        raise ExportError(
            f"Failed to export {fmt} bundle: {exc}", context={"output_path": str(output_path), "format": fmt}
        ) from exc
    return output_path


def _export_pdf(
    dossier: Dossier,
    output_path: Path,
    *,
    match_result: MatchResult | None,
    narrative: str | None,
    title: str,
) -> None:
    """Render the dossier bundle as a PDF using reportlab.

    Args:
        dossier: The parsed dossier to render.
        output_path: Destination PDF path.
        match_result: Optional match result to summarize.
        narrative: Optional narrative summary paragraph.
        title: Document title.
    """
    styles = getSampleStyleSheet()
    heading_style = styles["Heading2"]
    body_style = styles["BodyText"]
    title_style = ParagraphStyle("DossierTitle", parent=styles["Title"], spaceAfter=12)

    story: list[object] = [Paragraph(title, title_style)]

    if narrative:
        story.append(Paragraph("Summary", heading_style))
        story.append(Paragraph(narrative, body_style))
        story.append(Spacer(1, 0.2 * inch))

    if match_result is not None:
        story.append(Paragraph(f"G-Force Match Score: {match_result.g_force_score:.1f}%", heading_style))
        table_data = [["Category", "Weight", "Coverage"]]
        for score in match_result.category_scores:
            table_data.append([score.name, f"{score.weight:.2f}", f"{score.coverage * 100:.0f}%"])
        table = Table(table_data, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dddddd")),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                ]
            )
        )
        story.append(table)
        if match_result.gaps:
            story.append(Spacer(1, 0.1 * inch))
            story.append(Paragraph("Gaps: " + ", ".join(match_result.gaps), body_style))
        story.append(Spacer(1, 0.2 * inch))

    if dossier.capabilities:
        story.append(Paragraph("Capabilities", heading_style))
        by_domain: dict[str, list[str]] = {}
        for capability in dossier.capabilities:
            by_domain.setdefault(capability.domain, []).append(capability.text)
        for domain, items in by_domain.items():
            story.append(Paragraph(f"<b>{domain}</b>", body_style))
            for item in items:
                story.append(Paragraph(f"&bull; {item}", body_style))
        story.append(Spacer(1, 0.2 * inch))

    if dossier.chronology:
        story.append(Paragraph("Chronology", heading_style))
        for entry in dossier.chronology:
            story.append(Paragraph(f"<b>{entry.heading}</b>", body_style))
            for bullet in entry.bullets:
                story.append(Paragraph(f"&bull; {bullet}", body_style))
        story.append(Spacer(1, 0.2 * inch))

    if dossier.impact:
        story.append(Paragraph("Impact", heading_style))
        for stmt in dossier.impact:
            story.append(Paragraph(f"<b>{stmt.heading}</b>", body_style))
            for bullet in stmt.bullets:
                story.append(Paragraph(f"&bull; {bullet}", body_style))

    doc = SimpleDocTemplate(str(output_path), pagesize=LETTER)
    doc.build(story)


def _export_docx(
    dossier: Dossier,
    output_path: Path,
    *,
    match_result: MatchResult | None,
    narrative: str | None,
    title: str,
) -> None:
    """Render the dossier bundle as a DOCX using python-docx.

    Args:
        dossier: The parsed dossier to render.
        output_path: Destination DOCX path.
        match_result: Optional match result to summarize.
        narrative: Optional narrative summary paragraph.
        title: Document title.
    """
    document = Document()
    document.add_heading(title, level=0)

    if narrative:
        document.add_heading("Summary", level=1)
        document.add_paragraph(narrative)

    if match_result is not None:
        document.add_heading(f"G-Force Match Score: {match_result.g_force_score:.1f}%", level=1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid"
        header = table.rows[0].cells
        header[0].text, header[1].text, header[2].text = "Category", "Weight", "Coverage"
        for score in match_result.category_scores:
            row = table.add_row().cells
            row[0].text = score.name
            row[1].text = f"{score.weight:.2f}"
            row[2].text = f"{score.coverage * 100:.0f}%"
        if match_result.gaps:
            gaps_paragraph = document.add_paragraph()
            gaps_run = gaps_paragraph.add_run("Gaps: " + ", ".join(match_result.gaps))
            gaps_run.font.size = Pt(10)

    if dossier.capabilities:
        document.add_heading("Capabilities", level=1)
        by_domain: dict[str, list[str]] = {}
        for capability in dossier.capabilities:
            by_domain.setdefault(capability.domain, []).append(capability.text)
        for domain, items in by_domain.items():
            document.add_heading(domain, level=2)
            for item in items:
                document.add_paragraph(item, style="List Bullet")

    if dossier.chronology:
        document.add_heading("Chronology", level=1)
        for entry in dossier.chronology:
            document.add_heading(entry.heading, level=2)
            for bullet in entry.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    if dossier.impact:
        document.add_heading("Impact", level=1)
        for stmt in dossier.impact:
            document.add_heading(stmt.heading, level=2)
            for bullet in stmt.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    document.save(str(output_path))
