"""Unit tests for questvector.core.exporter."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from questvector.core.exporter import export_bundle
from questvector.core.models import (
    AlertLevel,
    Capability,
    CategoryScore,
    ChronologyEntry,
    Dossier,
    ImpactStatement,
    MatchResult,
)
from questvector.exceptions import ExportError


@pytest.fixture()
def full_dossier() -> Dossier:
    return Dossier(
        capabilities=(Capability(domain="Cloud", text="AWS and Terraform"),),
        chronology=(ChronologyEntry(heading="Acme — SRE", bullets=("Reduced cost 74%",)),),
        impact=(ImpactStatement(heading="Impact", bullets=("Saved $240K/yr",)),),
    )


@pytest.fixture()
def sample_match_result() -> MatchResult:
    return MatchResult(
        g_force_score=92.4,
        category_scores=(CategoryScore(name="cloud", weight=1.0, coverage=0.924, matched_keywords=("aws",)),),
        gaps=("kubernetes",),
        red_flags=(),
        alert_level=AlertLevel.SUCCESS,
    )


def test_export_pdf_produces_valid_pdf_file(tmp_path: Path, full_dossier: Dossier, sample_match_result: MatchResult) -> None:
    output_path = tmp_path / "bundle.pdf"
    export_bundle(full_dossier, output_path, fmt="pdf", match_result=sample_match_result, narrative="A summary.")

    assert output_path.exists()
    assert output_path.read_bytes()[:5] == b"%PDF-"


def test_export_docx_produces_reopenable_document(tmp_path: Path, full_dossier: Dossier, sample_match_result: MatchResult) -> None:
    output_path = tmp_path / "bundle.docx"
    export_bundle(full_dossier, output_path, fmt="docx", match_result=sample_match_result)

    assert output_path.exists()
    document = docx.Document(str(output_path))
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "Acme — SRE" in all_text or any("Acme" in h.text for h in document.paragraphs)


def test_export_handles_missing_optional_sections(tmp_path: Path) -> None:
    minimal_dossier = Dossier(capabilities=(Capability(domain="Cloud", text="AWS"),))
    output_path = tmp_path / "bundle.pdf"

    export_bundle(minimal_dossier, output_path, fmt="pdf")

    assert output_path.exists()


def test_export_unsupported_format_raises_export_error(tmp_path: Path, full_dossier: Dossier) -> None:
    with pytest.raises(ExportError):
        export_bundle(full_dossier, tmp_path / "bundle.txt", fmt="txt")  # type: ignore[arg-type]


def test_export_pdf_without_match_result_or_narrative(tmp_path: Path, full_dossier: Dossier) -> None:
    output_path = tmp_path / "bundle.pdf"
    export_bundle(full_dossier, output_path, fmt="pdf")
    assert output_path.read_bytes()[:5] == b"%PDF-"
